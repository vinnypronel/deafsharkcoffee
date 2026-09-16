from copy import deepcopy
from pathlib import Path
import shutil
import zipfile

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
LEGAL = ROOT / "legal"
FINAL = LEGAL / "SEND_TO_TECH_GUY"
ZIP_PATH = LEGAL / "SEND_TO_TECH_GUY_4_FILES.zip"


def merge_documents(base_path: Path, append_path: Path, output_path: Path, title: str, subtitle: str) -> None:
    doc = Document(base_path)

    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Title":
            paragraph.text = title
            break

    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Normal" and paragraph.text.strip():
            paragraph.text = subtitle
            break

    doc.add_page_break()
    appendix = Document(append_path)
    for element in appendix.element.body:
        if element.tag == qn("w:sectPr"):
            continue
        doc.element.body.insert(-1, deepcopy(element))

    doc.core_properties.title = title
    doc.core_properties.subject = subtitle
    doc.save(output_path)


def build() -> None:
    FINAL.mkdir(parents=True, exist_ok=True)

    outputs = [
        FINAL / "01_READ_FIRST_Website_Completion_Instructions.docx",
        FINAL / "02_WEBSITE_POLICIES_Terms_and_Privacy.docx",
        FINAL / "03_INTERNAL_PROCEDURES_Retention_and_Breach.docx",
        FINAL / "04_COPY_THIS_Email_to_Tech_Guy.txt",
    ]

    for existing in FINAL.iterdir():
        if existing.is_file() and existing not in outputs:
            existing.unlink()

    shutil.copy2(
        LEGAL / "Deaf_Shark_Coffee_Professional_Handoff_Package.docx",
        outputs[0],
    )

    merge_documents(
        LEGAL / "Deaf_Shark_Coffee_Privacy_Policy_Revised.docx",
        LEGAL / "Deaf_Shark_Coffee_Terms_of_Service_Revised.docx",
        outputs[1],
        "Website Terms and Privacy Policies",
        "Deaf Shark Coffee  |  Privacy Policy followed by Terms of Service  |  Attorney review copy",
    )

    merge_documents(
        LEGAL / "Deaf_Shark_Coffee_Data_Retention_and_Deletion_Procedure.docx",
        LEGAL / "Deaf_Shark_Coffee_Data_Breach_Response_Plan.docx",
        outputs[2],
        "Internal Data and Incident Procedures",
        "Deaf Shark Coffee  |  Retention and deletion procedure followed by breach response plan",
    )

    shutil.copy2(LEGAL / "04_COPY_THIS_Email_to_Tech_Guy.txt", outputs[3])

    with zipfile.ZipFile(ZIP_PATH, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for output in outputs:
            archive.write(output, output.name)

    print(f"Created {ZIP_PATH}")
    for output in outputs:
        print(output.name)


if __name__ == "__main__":
    build()
